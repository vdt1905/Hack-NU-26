"""
FormatForge AI — Phase 2 Test Suite
Tests the Transformation Engine (Agent 5) end-to-end.

Coverage:
 • Page layout (margins, page size)
 • Default typography (Normal style font, size, spacing)
 • Per-element formatting (title, author, abstract, headings, body, references)
 • Running head + page numbers (header content)w
 • Structural insertion (Abstract label, References label)
 • Page breaks (title-page, abstract-page, body-page, references-page)
 • Font enforcement on every run
 • Integration test on RS_TEST1.docx
"""

import os
import sys
import tempfile
import unittest
from pathlib import Path

# Ensure project root is on path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from backend.schemas.docir import (
    DocElement,
    DocIR,
    DocMetadata,
    ElementRole,
    ElementType,
    ParagraphFormatting,
    RunFormatting,
)
from backend.schemas.reports import ChangeRecord, ChangeStatus
from backend.schemas.style_spec import StyleSpec
from backend.agents.transformer import TransformerAgent


# ══════════════════════════════════════════════════════════════
#  Helpers
# ══════════════════════════════════════════════════════════════

def _make_elem(
    idx: int,
    role: ElementRole,
    content: str = "",
    bold: bool = False,
    italic: bool = False,
    font: str = "Calibri",
    size: float = 11.0,
) -> DocElement:
    """Create a minimal DocElement."""
    return DocElement(
        id=f"elem_{idx:04d}",
        type=ElementType.PARAGRAPH,
        role=role,
        content=content,
        formatting=ParagraphFormatting(
            font_name=font, font_size_pt=size, bold=bold, italic=italic,
        ),
        runs=[RunFormatting(text=content, bold=bold, italic=italic,
                            font_name=font, font_size_pt=size)],
    )


def _build_simple_docir() -> DocIR:
    """
    Build a small synthetic DocIR that mimics a minimal research paper:
      title → author → abstract_body → heading_1 → body(x3) → heading_1 → body → ref(x3)
    No ABSTRACT_LABEL or REFERENCE_LABEL — transformer should insert them.
    """
    elems = [
        _make_elem(1, ElementRole.TITLE, "The Effects of Sleep on Memory Consolidation", bold=True),
        _make_elem(2, ElementRole.AUTHOR_INFO, "John Smith"),
        _make_elem(3, ElementRole.AUTHOR_INFO, "University of Testing"),
        _make_elem(4, ElementRole.ABSTRACT_BODY, "This study examines the relationship between sleep duration and memory performance in college students."),
        _make_elem(5, ElementRole.ABSTRACT_BODY, "Results indicated significant positive correlations between sleep quality and recall."),
        _make_elem(6, ElementRole.KEYWORDS, "Keywords: sleep, memory, consolidation"),
        _make_elem(7, ElementRole.HEADING_1, "Introduction"),
        _make_elem(8, ElementRole.BODY, "Memory consolidation is a crucial cognitive process."),
        _make_elem(9, ElementRole.BODY, "Multiple studies have examined this relationship (Smith, 2020)."),
        _make_elem(10, ElementRole.BODY, "The current study extends this literature by testing new methods."),
        _make_elem(11, ElementRole.HEADING_1, "Results"),
        _make_elem(12, ElementRole.BODY, "Participants who slept 8+ hours recalled more items."),
        _make_elem(13, ElementRole.REFERENCE_ENTRY, "Smith, J. (2020). Sleep and memory. Journal of Cognitive Science, 15(3), 45-60."),
        _make_elem(14, ElementRole.REFERENCE_ENTRY, "Jones, A., & Williams, B. (2019). Memory consolidation during REM. Neuroscience, 22, 101-115."),
        _make_elem(15, ElementRole.REFERENCE_ENTRY, "Brown, C. (2018). Cognitive effects of sleep deprivation. Sleep Research, 10(1), 30-42."),
    ]
    return DocIR(
        metadata=DocMetadata(
            source_filename="test_paper.docx",
            source_format="docx",
            total_paragraphs=len(elems),
        ),
        elements=elems,
    )


def _create_docx_from_docir(docir: DocIR, path: Path) -> Path:
    """Create a real DOCX file matching the DocIR elements (one paragraph per element)."""
    doc = Document()
    for elem in docir.elements:
        if elem.type == ElementType.PARAGRAPH:
            para = doc.add_paragraph(elem.content)
            # Apply original formatting from the DocIR
            for run in para.runs:
                if elem.formatting.bold:
                    run.bold = True
                if elem.formatting.italic:
                    run.italic = True
                if elem.formatting.font_name:
                    run.font.name = elem.formatting.font_name
                if elem.formatting.font_size_pt:
                    run.font.size = Pt(elem.formatting.font_size_pt)
    doc.save(str(path))
    return path


def _run_transform(docir: DocIR, spec: StyleSpec = None) -> tuple[Path, list[ChangeRecord], Document]:
    """Helper: create DOCX, transform, return (output_path, changes, output_doc)."""
    if spec is None:
        spec = StyleSpec()  # APA 7 defaults
    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir = Path(tmpdir)
        input_path = tmpdir / "input.docx"
        _create_docx_from_docir(docir, input_path)
        agent = TransformerAgent()
        out_path, changes = agent.transform(docir, spec, input_path, output_dir=tmpdir)
        out_doc = Document(str(out_path))
        return out_path, changes, out_doc


# ══════════════════════════════════════════════════════════════
#  Test: Page Layout
# ══════════════════════════════════════════════════════════════

class TestPageLayout(unittest.TestCase):
    """Verify page margins, size, and orientation."""

    def setUp(self):
        self.docir = _build_simple_docir()
        _, self.changes, self.doc = _run_transform(self.docir)

    def test_margins_are_one_inch(self):
        for section in self.doc.sections:
            self.assertAlmostEqual(section.top_margin / 914400, 1.0, places=1)
            self.assertAlmostEqual(section.bottom_margin / 914400, 1.0, places=1)
            self.assertAlmostEqual(section.left_margin / 914400, 1.0, places=1)
            self.assertAlmostEqual(section.right_margin / 914400, 1.0, places=1)

    def test_page_size_letter(self):
        s = self.doc.sections[0]
        self.assertAlmostEqual(s.page_width / 914400, 8.5, places=1)
        self.assertAlmostEqual(s.page_height / 914400, 11.0, places=1)

    def test_page_layout_change_recorded(self):
        cats = [c.category for c in self.changes]
        self.assertIn("page_layout", cats)


# ══════════════════════════════════════════════════════════════
#  Test: Default Typography (Normal style)
# ══════════════════════════════════════════════════════════════

class TestDefaultTypography(unittest.TestCase):
    """Verify the Normal style is set correctly."""

    def setUp(self):
        self.docir = _build_simple_docir()
        _, self.changes, self.doc = _run_transform(self.docir)
        self.normal = self.doc.styles["Normal"]

    def test_font_name_times_new_roman(self):
        self.assertEqual(self.normal.font.name, "Times New Roman")

    def test_font_size_12pt(self):
        actual_pt = self.normal.font.size / 12700
        self.assertAlmostEqual(actual_pt, 12.0, places=0)

    def test_line_spacing_double(self):
        self.assertAlmostEqual(
            self.normal.paragraph_format.line_spacing, 2.0, places=1,
        )

    def test_space_after_zero(self):
        sa = self.normal.paragraph_format.space_after
        self.assertIsNotNone(sa)
        self.assertEqual(sa, Pt(0))

    def test_typography_change_recorded(self):
        cats = [c.category for c in self.changes]
        self.assertIn("typography", cats)


# ══════════════════════════════════════════════════════════════
#  Test: Title Formatting
# ══════════════════════════════════════════════════════════════

class TestTitleFormatting(unittest.TestCase):
    """Title paragraphs should be centered, bold, 12pt, no indent."""

    def setUp(self):
        self.docir = _build_simple_docir()
        _, self.changes, self.doc = _run_transform(self.docir)
        # First paragraph in output is the title (index 0 — but note: new
        # paragraphs may have been inserted.  Find the one with title text.)
        self.title_para = None
        for p in self.doc.paragraphs:
            if "Effects of Sleep" in p.text:
                self.title_para = p
                break

    def test_title_centered(self):
        self.assertIsNotNone(self.title_para)
        self.assertEqual(self.title_para.alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_title_bold(self):
        for run in self.title_para.runs:
            if run.text.strip():
                self.assertTrue(run.bold)

    def test_title_font(self):
        for run in self.title_para.runs:
            if run.text.strip():
                self.assertEqual(run.font.name, "Times New Roman")

    def test_title_no_indent(self):
        fi = self.title_para.paragraph_format.first_line_indent
        self.assertTrue(fi is None or fi == 0 or fi == Inches(0))

    def test_title_change_recorded(self):
        cats = [c.category for c in self.changes]
        self.assertIn("title_page", cats)


# ══════════════════════════════════════════════════════════════
#  Test: Author Info Formatting
# ══════════════════════════════════════════════════════════════

class TestAuthorInfoFormatting(unittest.TestCase):
    """Author lines should be centered, not bold."""

    def setUp(self):
        self.docir = _build_simple_docir()
        _, _, self.doc = _run_transform(self.docir)
        self.author_paras = [
            p for p in self.doc.paragraphs
            if "John Smith" in p.text or "University of Testing" in p.text
        ]

    def test_author_centered(self):
        for p in self.author_paras:
            self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_author_not_bold(self):
        for p in self.author_paras:
            for run in p.runs:
                if run.text.strip():
                    self.assertFalse(run.bold)


# ══════════════════════════════════════════════════════════════
#  Test: Abstract Formatting
# ══════════════════════════════════════════════════════════════

class TestAbstractFormatting(unittest.TestCase):
    """Abstract body should have no indent; label should be inserted."""

    def setUp(self):
        self.docir = _build_simple_docir()
        _, self.changes, self.doc = _run_transform(self.docir)
        self.abstract_paras = [
            p for p in self.doc.paragraphs
            if "examines the relationship" in p.text or "significant positive" in p.text
        ]

    def test_abstract_no_indent(self):
        for p in self.abstract_paras:
            fi = p.paragraph_format.first_line_indent
            self.assertTrue(fi is None or fi == 0 or fi == Inches(0))

    def test_abstract_left_aligned(self):
        for p in self.abstract_paras:
            self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.LEFT)

    def test_abstract_label_inserted(self):
        """An 'Abstract' label paragraph should exist before abstract body."""
        texts = [p.text for p in self.doc.paragraphs]
        self.assertIn("Abstract", texts)

    def test_abstract_label_centered_bold(self):
        for p in self.doc.paragraphs:
            if p.text == "Abstract":
                self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.CENTER)
                for run in p.runs:
                    if run.text.strip():
                        self.assertTrue(run.bold)
                break

    def test_abstract_on_new_page(self):
        """The first abstract paragraph (or its label) should start a new page."""
        found = False
        for p in self.doc.paragraphs:
            if p.text == "Abstract":
                self.assertTrue(p.paragraph_format.page_break_before)
                found = True
                break
        self.assertTrue(found, "Abstract label not found in output")


# ══════════════════════════════════════════════════════════════
#  Test: Keywords Formatting
# ══════════════════════════════════════════════════════════════

class TestKeywordsFormatting(unittest.TestCase):
    """Keywords line should be indented with 'Keywords:' in italic."""

    def setUp(self):
        self.docir = _build_simple_docir()
        _, _, self.doc = _run_transform(self.docir)
        self.kw_para = None
        for p in self.doc.paragraphs:
            if "sleep, memory, consolidation" in p.text:
                self.kw_para = p
                break

    def test_keywords_indented(self):
        self.assertIsNotNone(self.kw_para)
        fi = self.kw_para.paragraph_format.first_line_indent
        self.assertIsNotNone(fi)
        self.assertAlmostEqual(fi / 914400, 0.5, places=1)

    def test_keywords_italic_label(self):
        """At least one run containing 'Keywords' should be italic."""
        found_italic = False
        for run in self.kw_para.runs:
            if "keyword" in run.text.lower():
                found_italic = run.italic
        self.assertTrue(found_italic)


# ══════════════════════════════════════════════════════════════
#  Test: Heading Formatting
# ══════════════════════════════════════════════════════════════

class TestHeadingFormatting(unittest.TestCase):
    """Headings should follow APA 5-level system."""

    def setUp(self):
        self.docir = _build_simple_docir()
        _, self.changes, self.doc = _run_transform(self.docir)
        self.intro_para = None
        self.results_para = None
        for p in self.doc.paragraphs:
            if p.text == "Introduction":
                self.intro_para = p
            if p.text == "Results":
                self.results_para = p

    def test_h1_centered(self):
        self.assertIsNotNone(self.intro_para)
        self.assertEqual(self.intro_para.alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_h1_bold(self):
        for run in self.intro_para.runs:
            if run.text.strip():
                self.assertTrue(run.bold)

    def test_h1_not_italic(self):
        for run in self.intro_para.runs:
            if run.text.strip():
                self.assertFalse(run.italic)

    def test_heading_font(self):
        for run in self.intro_para.runs:
            if run.text.strip():
                self.assertEqual(run.font.name, "Times New Roman")

    def test_heading_change_recorded(self):
        heading_changes = [c for c in self.changes if c.category == "headings"]
        self.assertGreaterEqual(len(heading_changes), 2)  # at least 2 headings


class TestHeadingLevels(unittest.TestCase):
    """Test all 5 heading levels individually."""

    def _make_heading_docir(self, levels: list[int]) -> DocIR:
        elems = [_make_elem(1, ElementRole.TITLE, "Test", bold=True)]
        for i, lev in enumerate(levels):
            role = ElementRole(f"heading_{lev}")
            elems.append(_make_elem(i + 2, role, f"Heading Level {lev}"))
            elems.append(_make_elem(i + 100, ElementRole.BODY, "Body text here."))
        return DocIR(
            metadata=DocMetadata(source_filename="test.docx", total_paragraphs=len(elems)),
            elements=elems,
        )

    def test_level_2_left_bold(self):
        docir = self._make_heading_docir([2])
        _, _, doc = _run_transform(docir)
        for p in doc.paragraphs:
            if "Heading Level 2" in p.text:
                self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.LEFT)
                for run in p.runs:
                    if run.text.strip():
                        self.assertTrue(run.bold)
                        self.assertFalse(run.italic)

    def test_level_3_left_bold_italic(self):
        docir = self._make_heading_docir([3])
        _, _, doc = _run_transform(docir)
        for p in doc.paragraphs:
            if "Heading Level 3" in p.text:
                self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.LEFT)
                for run in p.runs:
                    if run.text.strip():
                        self.assertTrue(run.bold)
                        self.assertTrue(run.italic)

    def test_level_4_indented_bold(self):
        docir = self._make_heading_docir([4])
        _, _, doc = _run_transform(docir)
        for p in doc.paragraphs:
            if "Heading Level 4" in p.text:
                self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.LEFT)
                fi = p.paragraph_format.first_line_indent
                self.assertIsNotNone(fi)
                self.assertAlmostEqual(fi / 914400, 0.5, places=1)
                for run in p.runs:
                    if run.text.strip():
                        self.assertTrue(run.bold)
                        self.assertFalse(run.italic)

    def test_level_5_indented_bold_italic(self):
        docir = self._make_heading_docir([5])
        _, _, doc = _run_transform(docir)
        for p in doc.paragraphs:
            if "Heading Level 5" in p.text:
                for run in p.runs:
                    if run.text.strip():
                        self.assertTrue(run.bold)
                        self.assertTrue(run.italic)


# ══════════════════════════════════════════════════════════════
#  Test: Body Text Formatting
# ══════════════════════════════════════════════════════════════

class TestBodyFormatting(unittest.TestCase):
    """Body paragraphs should be left-aligned, 0.5\" first indent, double-spaced."""

    def setUp(self):
        self.docir = _build_simple_docir()
        _, self.changes, self.doc = _run_transform(self.docir)
        self.body_paras = [
            p for p in self.doc.paragraphs
            if "crucial cognitive" in p.text or "extends this literature" in p.text
        ]

    def test_body_left_aligned(self):
        for p in self.body_paras:
            self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.LEFT)

    def test_body_first_line_indent(self):
        for p in self.body_paras:
            fi = p.paragraph_format.first_line_indent
            self.assertIsNotNone(fi)
            self.assertAlmostEqual(fi / 914400, 0.5, places=1)

    def test_body_double_spaced(self):
        for p in self.body_paras:
            ls = p.paragraph_format.line_spacing
            self.assertIsNotNone(ls)
            self.assertAlmostEqual(ls, 2.0, places=1)

    def test_body_font_enforced(self):
        for p in self.body_paras:
            for run in p.runs:
                if run.text.strip():
                    self.assertEqual(run.font.name, "Times New Roman")

    def test_body_change_summary_recorded(self):
        body_changes = [c for c in self.changes if c.category == "body"]
        self.assertTrue(len(body_changes) >= 1)


# ══════════════════════════════════════════════════════════════
#  Test: Reference Formatting
# ══════════════════════════════════════════════════════════════

class TestReferenceFormatting(unittest.TestCase):
    """References should have hanging indent and double spacing."""

    def setUp(self):
        self.docir = _build_simple_docir()
        _, self.changes, self.doc = _run_transform(self.docir)
        self.ref_paras = [
            p for p in self.doc.paragraphs
            if "Sleep and memory" in p.text or "Memory consolidation during" in p.text
        ]

    def test_ref_hanging_indent(self):
        for p in self.ref_paras:
            li = p.paragraph_format.left_indent
            fi = p.paragraph_format.first_line_indent
            self.assertIsNotNone(li)
            self.assertAlmostEqual(li / 914400, 0.5, places=1)
            self.assertIsNotNone(fi)
            self.assertAlmostEqual(fi / 914400, -0.5, places=1)

    def test_ref_double_spaced(self):
        for p in self.ref_paras:
            ls = p.paragraph_format.line_spacing
            self.assertAlmostEqual(ls, 2.0, places=1)

    def test_ref_not_bold(self):
        for p in self.ref_paras:
            for run in p.runs:
                if run.text.strip():
                    self.assertFalse(run.bold)

    def test_references_label_inserted(self):
        texts = [p.text for p in self.doc.paragraphs]
        self.assertIn("References", texts)

    def test_references_label_centered_bold(self):
        for p in self.doc.paragraphs:
            if p.text == "References":
                self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.CENTER)
                for run in p.runs:
                    if run.text.strip():
                        self.assertTrue(run.bold)
                break

    def test_references_on_new_page(self):
        for p in self.doc.paragraphs:
            if p.text == "References":
                self.assertTrue(p.paragraph_format.page_break_before)
                break

    def test_ref_change_summary(self):
        ref_changes = [c for c in self.changes if c.category == "references"]
        self.assertGreaterEqual(len(ref_changes), 1)


# ══════════════════════════════════════════════════════════════
#  Test: Running Head + Page Numbers
# ══════════════════════════════════════════════════════════════

class TestRunningHead(unittest.TestCase):
    """Header should contain running head text + PAGE field."""

    def setUp(self):
        self.docir = _build_simple_docir()
        spec = StyleSpec()
        # Enable running head for professional paper
        spec.running_head.enabled = True
        _, self.changes, self.doc = _run_transform(self.docir, spec)

    def test_header_has_text(self):
        header = self.doc.sections[0].header
        full_text = "".join(p.text for p in header.paragraphs)
        self.assertIn("EFFECTS OF SLEEP", full_text.upper())

    def test_header_has_page_field(self):
        """The header XML should contain a PAGE field code."""
        header = self.doc.sections[0].header
        xml_str = header._element.xml
        self.assertIn("PAGE", xml_str)

    def test_header_font(self):
        header = self.doc.sections[0].header
        for p in header.paragraphs:
            for run in p.runs:
                if run.font.name:
                    self.assertEqual(run.font.name, "Times New Roman")

    def test_running_head_change_recorded(self):
        cats = [c.category for c in self.changes]
        self.assertIn("running_head", cats)


class TestPageNumbersOnly(unittest.TestCase):
    """When running_head.enabled=False, still get page numbers."""

    def setUp(self):
        self.docir = _build_simple_docir()
        spec = StyleSpec()
        spec.running_head.enabled = False
        _, self.changes, self.doc = _run_transform(self.docir, spec)

    def test_page_field_present(self):
        header = self.doc.sections[0].header
        xml_str = header._element.xml
        self.assertIn("PAGE", xml_str)


# ══════════════════════════════════════════════════════════════
#  Test: Page Break Structure
# ══════════════════════════════════════════════════════════════

class TestPageBreaks(unittest.TestCase):
    """Verify APA page structure: title page → abstract → body."""

    def setUp(self):
        self.docir = _build_simple_docir()
        _, _, self.doc = _run_transform(self.docir)
        self.paragraphs = self.doc.paragraphs

    def test_abstract_starts_new_page(self):
        """'Abstract' label should have page_break_before."""
        for p in self.paragraphs:
            if p.text == "Abstract":
                self.assertTrue(p.paragraph_format.page_break_before)
                return
        self.fail("Abstract label not found")

    def test_body_starts_new_page(self):
        """First heading after abstract should start a new page."""
        for p in self.paragraphs:
            if p.text == "Introduction":
                self.assertTrue(p.paragraph_format.page_break_before)
                return
        self.fail("Introduction heading not found")


# ══════════════════════════════════════════════════════════════
#  Test: Font Enforcement
# ══════════════════════════════════════════════════════════════

class TestFontEnforcement(unittest.TestCase):
    """Every text-bearing run should be Times New Roman after transform."""

    def setUp(self):
        self.docir = _build_simple_docir()
        _, _, self.doc = _run_transform(self.docir)

    def test_all_body_runs_have_correct_font(self):
        wrong_fonts: list[str] = []
        for para in self.doc.paragraphs:
            for run in para.runs:
                if run.text.strip() and run.font.name and run.font.name != "Times New Roman":
                    wrong_fonts.append(f"{run.text[:30]}: {run.font.name}")
        self.assertEqual(wrong_fonts, [], f"Runs with wrong fonts: {wrong_fonts}")


# ══════════════════════════════════════════════════════════════
#  Test: Change Records
# ══════════════════════════════════════════════════════════════

class TestChangeRecords(unittest.TestCase):
    """Transform should produce meaningful change records."""

    def setUp(self):
        self.docir = _build_simple_docir()
        _, self.changes, _ = _run_transform(self.docir)

    def test_changes_not_empty(self):
        self.assertGreater(len(self.changes), 5)

    def test_all_applied(self):
        for c in self.changes:
            self.assertEqual(c.status, ChangeStatus.APPLIED)

    def test_category_diversity(self):
        cats = set(c.category for c in self.changes)
        # Should have at least: page_layout, typography, title_page, abstract,
        # headings, body, references, running_head
        self.assertGreaterEqual(len(cats), 5, f"Categories: {cats}")

    def test_rule_references_present(self):
        with_refs = [c for c in self.changes if c.rule_reference]
        self.assertGreater(len(with_refs), 3)


# ══════════════════════════════════════════════════════════════
#  Test: Edge Cases
# ══════════════════════════════════════════════════════════════

class TestEdgeCases(unittest.TestCase):
    """Edge cases: empty paragraphs, single element, no abstract, etc."""

    def test_empty_paragraph_no_crash(self):
        """An empty-content element should not crash the transformer."""
        elems = [
            _make_elem(1, ElementRole.TITLE, "A Title"),
            _make_elem(2, ElementRole.BODY, ""),
            _make_elem(3, ElementRole.BODY, "Some text."),
        ]
        docir = DocIR(
            metadata=DocMetadata(source_filename="t.docx", total_paragraphs=3),
            elements=elems,
        )
        _, changes, doc = _run_transform(docir)
        self.assertIsNotNone(doc)

    def test_no_abstract_no_crash(self):
        """Document with no abstract should transform without error."""
        elems = [
            _make_elem(1, ElementRole.TITLE, "Title Only"),
            _make_elem(2, ElementRole.BODY, "Body text."),
        ]
        docir = DocIR(
            metadata=DocMetadata(source_filename="t.docx", total_paragraphs=2),
            elements=elems,
        )
        _, changes, doc = _run_transform(docir)
        self.assertIsNotNone(doc)

    def test_no_references_no_crash(self):
        """Document with no references should transform without error."""
        elems = [
            _make_elem(1, ElementRole.TITLE, "Title"),
            _make_elem(2, ElementRole.ABSTRACT_BODY, "Abstract text."),
            _make_elem(3, ElementRole.BODY, "Body."),
        ]
        docir = DocIR(
            metadata=DocMetadata(source_filename="t.docx", total_paragraphs=3),
            elements=elems,
        )
        _, changes, doc = _run_transform(docir)
        self.assertIsNotNone(doc)

    def test_all_unknown_roles(self):
        """All UNKNOWN elements should get body treatment."""
        elems = [
            _make_elem(1, ElementRole.UNKNOWN, "Unknown content 1."),
            _make_elem(2, ElementRole.UNKNOWN, "Unknown content 2."),
        ]
        docir = DocIR(
            metadata=DocMetadata(source_filename="t.docx", total_paragraphs=2),
            elements=elems,
        )
        _, _, doc = _run_transform(docir)
        for p in doc.paragraphs:
            if p.text.strip():
                ls = p.paragraph_format.line_spacing
                self.assertAlmostEqual(ls, 2.0, places=1)

    def test_existing_abstract_label_not_duplicated(self):
        """If ABSTRACT_LABEL exists, don't insert a second one."""
        elems = [
            _make_elem(1, ElementRole.TITLE, "Title"),
            _make_elem(2, ElementRole.ABSTRACT_LABEL, "Abstract"),
            _make_elem(3, ElementRole.ABSTRACT_BODY, "Body of abstract."),
            _make_elem(4, ElementRole.BODY, "Main body."),
        ]
        docir = DocIR(
            metadata=DocMetadata(source_filename="t.docx", total_paragraphs=4),
            elements=elems,
        )
        _, _, doc = _run_transform(docir)
        abstract_count = sum(1 for p in doc.paragraphs if p.text.strip() == "Abstract")
        self.assertEqual(abstract_count, 1)


# ══════════════════════════════════════════════════════════════
#  Test: Integration — RS_TEST1.docx
# ══════════════════════════════════════════════════════════════

class TestRS_TEST1_Integration(unittest.TestCase):
    """Full pipeline integration test on the real test document."""

    @classmethod
    def setUpClass(cls):
        test_doc = Path("RS_TEST1.docx")
        if not test_doc.exists():
            raise unittest.SkipTest("RS_TEST1.docx not found")

        from backend.agents.ingest import IngestAgent
        from backend.agents.structure_detector import StructureDetectorAgent
        from backend.agents.citation_engine import CitationEngineAgent

        # Run pipeline
        cls.docir = IngestAgent().parse(test_doc)
        StructureDetectorAgent().detect(cls.docir)
        spec = StyleSpec()
        CitationEngineAgent().process(cls.docir, spec)

        cls.spec = spec
        cls.agent = TransformerAgent()

        cls._tmpdir = tempfile.mkdtemp()
        tmpdir = Path(cls._tmpdir)
        cls.output_path, cls.changes = cls.agent.transform(
            cls.docir, spec, test_doc, output_dir=tmpdir,
        )
        cls.out_doc = Document(str(cls.output_path))

    @classmethod
    def tearDownClass(cls):
        import shutil
        if hasattr(cls, "_tmpdir") and os.path.exists(cls._tmpdir):
            shutil.rmtree(cls._tmpdir, ignore_errors=True)

    def test_output_file_created(self):
        self.assertTrue(self.output_path.exists())

    def test_margins_correct(self):
        s = self.out_doc.sections[0]
        self.assertAlmostEqual(s.top_margin / 914400, 1.0, places=1)

    def test_font_is_tnr(self):
        normal = self.out_doc.styles["Normal"]
        self.assertEqual(normal.font.name, "Times New Roman")

    def test_line_spacing_double(self):
        normal = self.out_doc.styles["Normal"]
        self.assertAlmostEqual(normal.paragraph_format.line_spacing, 2.0, places=1)

    def test_changes_recorded(self):
        self.assertGreater(len(self.changes), 10)

    def test_change_categories(self):
        cats = set(c.category for c in self.changes)
        self.assertIn("page_layout", cats)
        self.assertIn("typography", cats)
        self.assertIn("headings", cats)

    def test_body_count_in_changes(self):
        body_changes = [c for c in self.changes if c.category == "body"]
        self.assertGreaterEqual(len(body_changes), 1)

    def test_header_has_page_number(self):
        header = self.out_doc.sections[0].header
        xml_str = header._element.xml
        self.assertIn("PAGE", xml_str)

    def test_abstract_label_present(self):
        """After transform, 'Abstract' text should appear in the document."""
        found = any(p.text.strip() == "Abstract" for p in self.out_doc.paragraphs)
        self.assertTrue(found, "Abstract label not inserted")

    def test_references_label_present(self):
        """After transform, 'References' text should appear in the document."""
        found = any(p.text.strip() == "References" for p in self.out_doc.paragraphs)
        self.assertTrue(found, "References label not inserted")


# ══════════════════════════════════════════════════════════════

if __name__ == "__main__":
    unittest.main()
